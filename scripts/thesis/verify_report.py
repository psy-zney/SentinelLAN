"""Check the audited regressions and render every final PDF page for visual QA.

Exit 1 means a document defect, exit 2 means a requested publication gate still
has missing acceptance evidence, and exit 0 means the structural checks pass.
Visual inspection remains necessary; this checker does not claim to recognize
all overlaps inside raster screenshots or prove real OS acceptance.
"""
from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import sys

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
import pdfplumber

from revision_content import REFERENCES

FORBIDDEN = ["scratch/", "hồ sơ kiểm thử được đọc", "không chạy lại các bộ test",
             "Chưa thu ảnh", "Chưa nghiệm thu", "(tài liệu dự án A", "mutation",
             "denominator", "hợp instrument", "endpoint máy của tôi", "Error!", "Lỗi!"]


def xml_text(element):
    return "".join(element.xpath(".//w:t/text()"))


def check_report(docx: Path, pdf: Path, metadata: Path | None, qa: Path, render=False, publication=False):
    errors, warnings = [], []
    document = Document(docx)
    pdf_reader = PdfReader(pdf)
    paragraph_text = [xml_text(p._p) for p in document.paragraphs]
    table_text = [c.text for t in document.tables for row in t.rows for c in row.cells]
    all_text = "\n".join(paragraph_text+table_text)
    clean_text = all_text.replace("\u200b", "")
    # 1. Check every explicit font family, including hidden field runs and styles.
    fonts = set()
    for part in document.part.package.parts:
        root = getattr(part, "element", None)
        if root is None:
            continue
        for font in root.iter(qn("w:rFonts")):
            for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
                value = font.get(qn("w:"+slot))
                if value and value != "Times New Roman":
                    errors.append(f"DOCX font {slot}: {value}")
    for page in pdf_reader.pages:
        for entry in page["/Resources"].get("/Font", {}).values():
            family = str(entry.get_object().get("/BaseFont", ""))
            fonts.add(family)
            if "TimesNewRomanPS" not in family:
                errors.append("PDF font: "+family)
    # 2. Cover data, verified independently of the audit's unconfirmed assertions.
    cover = pdf_reader.pages[0].extract_text() or ""
    for value in ["TRƯỜNG CÔNG NGHỆ VÀ THIẾT KẾ", "TS. Nguyễn An Tế", "31231025401", "Công nghệ thông tin"]:
        if value not in cover:
            errors.append("Cover missing: "+value)
    if re.search(r"(?:GVHD|MSSV) +:", cover):
        errors.append("Extra spaces before a cover colon")
    # 3. Styles and rendered left edge of all chapter-level TOC entries.
    for style in document.styles:
        name = style.name.lower()
        if name in ("toc 1", "table of figures"):
            left = style.paragraph_format.left_indent
            if left is not None and abs(left.cm) > .03:
                errors.append("Incorrect left indent in "+name)
        if name == "toc 1":
            effective_size = style.font.size
            parent = style.base_style
            while effective_size is None and parent:
                effective_size = parent.font.size
                parent = parent.base_style
            if not style.font.bold or effective_size is None or abs(effective_size.pt-13) > .1:
                errors.append("TOC level 1 must be bold 13 pt")
    pages, toc_lines, fragment_pages = [], [], []
    with pdfplumber.open(pdf) as report_pdf:
        body_started = False
        for i, page in enumerate(report_pdf.pages):
            text = page.extract_text() or ""
            lines = [line for line in text.splitlines() if line.strip()]
            is_contents = "MỤC LỤC" in text or (i < 8 and "CHƯƠNG 1" in text and len(re.findall(r"CHƯƠNG [1-7]", text)) > 1)
            if is_contents:
                words = page.extract_words(extra_attrs=["fontname", "size"])
                for word in words:
                    if word["text"] == "CHƯƠNG":
                        toc_lines.append({"page": i+1, "x0": word["x0"], "font": word["fontname"], "size": word["size"]})
                        if abs(word["x0"]-3.5*72/2.54) > .05*72/2.54:
                            errors.append(f"TOC chapter left edge on page {i+1}: {word['x0']:.2f} pt")
                        if "Bold" not in word["fontname"]:
                            errors.append(f"TOC chapter is not bold on page {i+1}")
            if i > 6 and "CHƯƠNG 1" in text and not is_contents:
                body_started = True
            fragments = [line for line in lines if line.strip() in {"dAt", "ngs", "ctKey"}]
            if fragments:
                fragment_pages.append(i+1)
            has_graphic = bool(page.images or page.lines or page.rects)
            content_lines = [line for line in lines if not re.fullmatch(r"\d+|[ivxlcdm]+", line.strip())]
            if body_started and len(content_lines) < 8 and not has_graphic:
                errors.append(f"Underfilled content page {i+1}: {len(content_lines)} lines")
            pages.append({"page": i+1, "lines": len(content_lines), "graphics": has_graphic,
                          "start": text[:100].replace("\n", " ")})
    # 4. Dictionary widths and rendered broken suffixes.
    for t in document.tables:
        if t.rows[0].cells[0].text == "Model hoặc trường":
            widths = [c.width.cm for c in t.columns]
            if any(abs(a-b) > .04 for a, b in zip(widths, [6.3, 9.2])):
                errors.append(f"Data dictionary widths: {widths}")
    if fragment_pages:
        errors.append("Broken identifier suffixes on pages: "+str(fragment_pages))
    # 5. Actual print size, DPI, and text intersections in the generated diagrams.
    diagram_meta = docx.parent / "hinh-anh/figures.json"
    if diagram_meta.exists():
        embedded_images = {hashlib.sha256(part.blob).hexdigest()
                           for part in document.part.package.parts if part.content_type.startswith("image/")}
        for figure in json.loads(diagram_meta.read_text(encoding="utf-8")):
            figure_file = diagram_meta.parent / (figure["name"]+".png")
            if hashlib.sha256(figure_file.read_bytes()).hexdigest() not in embedded_images:
                errors.append("Revised diagram missing from DOCX: "+figure["name"])
            if figure["dpi"] < 300 or min(t["font_pt"] for t in figure["text"]) < 10:
                errors.append("Unreadable figure: "+figure["name"])
            for j, a in enumerate(figure["text"]):
                x1, y1, x2, y2 = a["bounds"]
                if x1 < 0 or x2 > figure["width_px"]:
                    errors.append("Clipped figure label: "+a["text"])
                for b in figure["text"][j+1:]:
                    bx1, by1, bx2, by2 = b["bounds"]
                    if min(x2, bx2) > max(x1, bx1) and min(y2, by2) > max(y1, by1):
                        errors.append("Overlapping figure labels: "+figure["name"])
    else:
        errors.append("Missing diagram size metadata")
    # An SVG alternative is allowed only on cover furniture, never on revised diagrams.
    for p in document.paragraphs:
        description = " ".join(p._p.xpath(".//wp:docPr/@descr"))
        if description.startswith("Sơ đồ đã sửa") and p._p.xpath(".//*[local-name()='svgBlip']"):
            errors.append("Stale SVG alternative in a revised diagram")
    # 6–7. Editorial residue and mixed-language defects in the final text.
    for token in FORBIDDEN:
        if token.lower() in clean_text.lower():
            errors.append("Forbidden text: "+token)
    # 8. Coverage of the previously omitted modules in the actual technical body.
    start = next((i for i, p in enumerate(document.paragraphs) if p.style.name == "Heading 1" and p.text.startswith("CHƯƠNG 3")), len(paragraph_text))
    end = next((i for i, p in enumerate(document.paragraphs) if p.style.name == "Heading 1" and p.text.startswith("CHƯƠNG 6")), start)
    technical = "\n".join(paragraph_text[start:end])
    for value in ["DeviceQrLabel", "WorkOrder", "AssetLoan", "Incident", "InstallApp", "Privilege", "Panic", "PauseAgent", "UninstallAgent", "Appointment", "PBKDF2", "600.000", "AES-256-GCM", "purpose", "SetKernelObjectSecurity", "WindowsDesktopBridge", "Session 0", "IsolationScript", "SYSTEM", "ApprovedAppInstaller"]:
        if value not in technical:
            errors.append("Technical body missing: "+value)
    # 9. Every bibliography entry has a matching author/year citation in the body.
    ref_index = next((i for i, p in enumerate(document.paragraphs) if p.style.name == "Heading 1" and p.text == "TÀI LIỆU THAM KHẢO"), len(paragraph_text))
    body = "\n".join(paragraph_text[:ref_index])
    for author, year, _, _ in REFERENCES:
        key = author.split(",")[0]
        if not re.search(re.escape(key)+r"[^\n]{0,110}?"+re.escape(year), body):
            errors.append(f"Uncited reference: {key}, {year}")
    if "google.com/search" in clean_text:
        errors.append("Search URL used as a bibliography source")
    # Native captions must carry chapter prefixes and sequence fields.
    for p in document.paragraphs:
        if p.style.name == "Caption":
            text = xml_text(p._p)
            if not re.match(r"(?:Bảng|Hình) [1-7A-E]\.\d+: ", text):
                errors.append("Caption numbering: "+text)
    # 10. Scenario evidence: do not turn partial/OS checks into a full acceptance pass.
    evidence = json.loads(metadata.read_text(encoding="utf-8")) if metadata and metadata.exists() else None
    if evidence:
        if evidence["dotnet"] != {"Passed": 380, "NotExecuted": 2}:
            errors.append("Incorrect .NET pass/skip accounting")
        for scenario in evidence["scenarios"]:
            if scenario["status"].startswith("Đạt") and (not scenario["evidence"] or any(e["outcome"] != "Passed" for e in scenario["evidence"])):
                errors.append("Unsupported scenario result: "+scenario["code"])
        warnings.extend(evidence["publication_blockers"])
        if len(evidence["screenshots"]) < 6:
            warnings.append("Fewer than six verified product screenshots")
    else:
        errors.append("Missing scenario evidence manifest")
    qa.mkdir(parents=True, exist_ok=True)
    if render:
        import pypdfium2 as pdfium
        from PIL import Image, ImageDraw, ImageFont
        image_dir = qa / "pages"
        image_dir.mkdir(exist_ok=True)
        pdf_document = pdfium.PdfDocument(pdf)
        for i, page in enumerate(pdf_document):
            page.render(scale=1.7).to_pil().save(image_dir / f"page-{i+1:02}.png")
            page.close()
        for offset in range(0, len(pages), 6):
            sheet = Image.new("RGB", (1500, 1500), "#cccccc")
            draw = ImageDraw.Draw(sheet)
            font = ImageFont.truetype("C:/Windows/Fonts/times.ttf", 24)
            for j in range(offset, min(offset+6, len(pages))):
                im = Image.open(image_dir / f"page-{j+1:02}.png")
                im.thumbnail((480, 690))
                x, y = (j-offset)%3*500+10, (j-offset)//3*750+40
                sheet.paste(im, (x, y))
                draw.text((x, y-30), f"Trang {j+1}", fill="black", font=font)
            sheet.save(qa / f"sheet-{offset//6+1:02}.jpg", quality=92)
    result = {"errors": sorted(set(errors)), "publication_blockers": sorted(set(warnings)),
              "pdf_fonts": sorted(fonts), "page_count": len(pages), "pages": pages, "toc_chapter_lines": toc_lines,
              "limitations": ["Diagram connector routing and final page readability require visual review.",
                              "Automated fixture passes do not prove real OS acceptance or long-term performance."]}
    (qa / "verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: result[k] for k in ["errors", "publication_blockers", "page_count"]}, ensure_ascii=False))
    return 1 if errors else 2 if publication and warnings else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--docx", required=True, type=Path)
    parser.add_argument("--pdf", required=True, type=Path)
    parser.add_argument("--metadata", type=Path)
    parser.add_argument("--qa", required=True, type=Path)
    parser.add_argument("--render", action="store_true")
    parser.add_argument("--publication", action="store_true")
    args = parser.parse_args()
    raise SystemExit(check_report(args.docx, args.pdf, args.metadata, args.qa, args.render, args.publication))
