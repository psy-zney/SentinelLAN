import sys
import re
import json
from pathlib import Path
from collections import defaultdict
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / 'Khoa-luan-SentinelLAN-ban-thao.md'
OUT = ROOT / 'Khoa-luan-SentinelLAN.docx'
RED = RGBColor.from_string('9C0006')
BLACK = RGBColor(0, 0, 0)
DOC = Document()
FIGURES = json.loads((ROOT / 'figure-plan.json').read_text(encoding='utf-8'))
if isinstance(FIGURES, dict):
    FIGURES = FIGURES.get('figures', FIGURES.get('items', []))
FIG = {x['id']: x for x in FIGURES}
assert len(FIG) == 22
TABLES = []
CURRENT_CHAPTER = 0
TABLE_COUNTS = defaultdict(int)
FIGURE_INSERTED = []
BOOKMARK = 1


def font(style, size, bold=False, italic=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = BLACK
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rfonts.set(qn('w:' + key), 'Times New Roman')
    for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme', 'csTheme']:
        rfonts.attrib.pop(qn('w:' + key), None)


def setup():
    s = DOC.sections[0]
    s.page_width = Cm(21)
    s.page_height = Cm(29.7)
    s.top_margin = Cm(2.5)
    s.bottom_margin = Cm(2.5)
    s.left_margin = Cm(3)
    s.right_margin = Cm(2)
    s.header_distance = Cm(1.2)
    s.footer_distance = Cm(1.2)
    s.different_first_page_header_footer = True
    normal = DOC.styles['Normal']
    font(normal, 13)
    pf = normal.paragraph_format
    pf.line_spacing = 1.3
    pf.space_after = Pt(6)
    pf.first_line_indent = Cm(0.7)
    pf.widow_control = True
    for name, size in [('Title', 18), ('Subtitle', 13), ('Heading 1', 16), ('Heading 2', 14), ('Heading 3', 13)]:
        st = DOC.styles[name]
        font(st, size, name != 'Subtitle')
        st.paragraph_format.first_line_indent = Cm(0)
        st.paragraph_format.space_before = Pt(14 if name != 'Heading 3' else 10)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.line_spacing = 1.12
        st.paragraph_format.keep_with_next = True
    for name, size, italic in [('Ghi chu bien tap', 11.5, True), ('Chu thich hinh', 11.5, True), ('Chu thich bang', 11.5, False), ('Table Text', 11, False), ('Front Heading', 16, False), ('Figure Guide Heading', 13, False)]:
        st = DOC.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = normal
        font(st, size, name == 'Front Heading', italic)
        st.paragraph_format.first_line_indent = Cm(0)
        st.paragraph_format.line_spacing = 1.1
        st.paragraph_format.space_after = Pt(6)
        if name in ['Chu thich hinh', 'Chu thich bang', 'Front Heading']:
            st.paragraph_format.keep_with_next = name != 'Chu thich hinh'
        if name in ['Ghi chu bien tap','Figure Guide Heading']:
            st.paragraph_format.keep_together = True
        if name == 'Figure Guide Heading':
            st.paragraph_format.keep_with_next = True
            st.font.bold = True
    DOC.styles['Ghi chu bien tap'].font.color.rgb = RED
    for i in range(1,4):
        name = f'TOC {i}'
        st = DOC.styles[name] if name in DOC.styles else DOC.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        font(st, 12)
        st.paragraph_format.first_line_indent = Cm(0)
        st.paragraph_format.line_spacing = 1.08
        st.paragraph_format.space_after = Pt(3)
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    field(footer, 'PAGE', '1')
    DOC.core_properties.title = 'Hệ thống quản lý giám sát và bảo mật thiết bị đầu cuối trong mạng nội bộ'
    DOC.core_properties.subject = 'Khóa luận tốt nghiệp về hệ thống SentinelLAN'
    DOC.core_properties.author = 'Lê Quang Khánh'
    DOC.core_properties.keywords = 'SentinelLAN, khóa luận, quản lý thiết bị, LAN'
    DOC.core_properties.comments = ''
    upd = OxmlElement('w:updateFields')
    upd.set(qn('w:val'), 'true')
    DOC.settings.element.append(upd)
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), 'vi-VN')
    DOC.styles['Normal'].element.get_or_add_rPr().append(lang)
    for st in DOC.styles:
        for el in list(st.element.iter(qn('w:pBdr'))):
            el.getparent().remove(el)


def field(p, instruction, cached=''):
    r = p.add_run()
    a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'begin')
    b = OxmlElement('w:instrText'); b.set(qn('xml:space'), 'preserve'); b.text = ' ' + instruction + ' '
    c = OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'), 'separate')
    d = OxmlElement('w:t'); d.text = cached
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end')
    for x in [a,b,c,d,e]: r._r.append(x)
    return r


def bookmark(p, name):
    global BOOKMARK
    a = OxmlElement('w:bookmarkStart'); a.set(qn('w:id'), str(BOOKMARK)); a.set(qn('w:name'), name)
    b = OxmlElement('w:bookmarkEnd'); b.set(qn('w:id'), str(BOOKMARK))
    p._p.insert(0,a); p._p.append(b); BOOKMARK += 1


def rich(p, text):
    pattern = r'(\[(?:TỰ ĐIỀN|ĐO VÀ ĐIỀN|CHÈN HÌNH|CHỤP HÌNH|GHI CHÚ BIÊN TẬP|XÁC NHẬN|BỔ SUNG NGUỒN)[^\]]*\]|\*\*.*?\*\*|(?<!\*)\*[^*]+\*(?!\*))'
    for part in re.split(pattern, text):
        if not part: continue
        if re.match(r'^\[(TỰ ĐIỀN|ĐO VÀ ĐIỀN|CHÈN HÌNH|CHỤP HÌNH|GHI CHÚ BIÊN TẬP|XÁC NHẬN|BỔ SUNG NGUỒN)', part):
            r = p.add_run(part); r.font.color.rgb = RED; r.bold = True; r.italic = True
        elif part.startswith('**') and part.endswith('**'):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith('*') and part.endswith('*'):
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(part)


def para(text, style=None, keep=False):
    p = DOC.add_paragraph(style=style)
    p.paragraph_format.keep_with_next = keep
    if style != 'Ghi chu bien tap':
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, text)
    return p


def note(text, keep=False):
    p = para(text, 'Ghi chu bien tap', keep)
    return p


def front_heading(text):
    p = DOC.add_paragraph(text, style='Front Heading')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def cover():
    for text, size, bold, after in [
        ('ĐẠI HỌC KINH TẾ THÀNH PHỐ HỒ CHÍ MINH',14,True,10),
        ('KHOA CÔNG NGHỆ THÔNG TIN',14,True,48),
        ('KHÓA LUẬN TỐT NGHIỆP',18,True,34),
        ('HỆ THỐNG QUẢN LÝ GIÁM SÁT VÀ BẢO MẬT THIẾT BỊ ĐẦU CUỐI TRONG MẠNG NỘI BỘ',18,True,22),
        ('SentinelLAN',15,False,45),
    ]:
        p=DOC.add_paragraph(style='Title' if text.startswith('HỆ THỐNG') else 'Normal')
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.space_after=Pt(after)
        r=p.add_run(text); r.font.size=Pt(size); r.bold=bold
    for text in ['Sinh viên thực hiện: Lê Quang Khánh','Mã số sinh viên: [TỰ ĐIỀN A01: mã số sinh viên]','Lớp và khóa: [TỰ ĐIỀN A02: lớp, khóa học]','Ngành: Công nghệ thông tin','Giảng viên hướng dẫn: [TỰ ĐIỀN A03: học hàm, học vị, họ tên]']:
        p=para(text); p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.left_indent=Cm(1.7)
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p=para('TP. Hồ Chí Minh năm 2026')
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.space_before=Pt(36)
    DOC.add_page_break()


def editor_guide():
    front_heading('HƯỚNG DẪN HOÀN THIỆN FILE WORD')
    para('Tài liệu gồm phần nội dung khóa luận và các ghi chú giúp bổ sung minh chứng. Chữ màu đỏ là phần cần điền, xác nhận hoặc thay bằng dữ liệu thực tế. Bỏ trang hướng dẫn và các ghi chú biên tập sau khi hoàn thiện bản nộp.')
    for lead, body in [
        ('TỰ ĐIỀN','Thông tin cá nhân, bối cảnh khảo sát, cấu hình và môi trường thực nghiệm.'),
        ('ĐO VÀ ĐIỀN','Chạy kiểm thử hoặc đo trên hệ thống; điền số liệu thật và giữ bằng chứng. Chưa thực hiện phải được báo cáo đúng trạng thái.'),
        ('CHỤP ẢNH','Chụp đúng màn hình hoặc kết quả nêu tại vị trí hình. Che token, mật khẩu, khóa và thông tin riêng tư.'),
        ('VẼ SƠ ĐỒ','Dựng sơ đồ từ thiết kế hoặc phiên bản mã nguồn dùng bảo vệ; không chụp giao diện để thay sơ đồ.'),
        ('TẠO BIỂU ĐỒ','Tạo hình từ dữ liệu đo gốc; giữ đơn vị, nhãn trục và điều kiện đo.'),
        ('CHÚ THÍCH','Chú thích hình và bảng đã được đánh số theo chương. Khi thêm hoặc xóa hình, cập nhật trường trong Word để làm mới danh mục.'),
    ]:
        note(f'{lead}: {body}')
    para('Định dạng sử dụng khổ A4, Times New Roman, cỡ chữ nội dung 13, giãn dòng 1,3; lề trái 3 cm, phải 2 cm, trên và dưới 2,5 cm. Đối chiếu với mẫu chính thức của khoa trước khi nộp.')
    note('[TỰ ĐIỀN A04: xác nhận tên đơn vị, mẫu bìa, kiểu trích dẫn và quy định định dạng của trường.]')
    note('Kết quả kiểm thử và khảo sát chưa được thay bằng số liệu tự tạo. Chốt Chương 6 trước khi hoàn thiện tóm tắt và Chương 7. Sơ đồ kiến trúc tổng thể được chèn sẵn; các hình minh chứng còn lại có hướng dẫn ngay tại vị trí chèn.')
    para('Để cập nhật mục lục và danh mục: chọn toàn bộ nội dung bằng Ctrl+A, nhấn F9 và chọn cập nhật toàn bộ bảng mục lục khi Word hỏi. Nếu dùng bàn phím laptop, có thể cần Fn+F9. Các vị trí hình H01–H22 được tổng hợp tại Phụ lục G.')
    DOC.add_page_break()


def lists():
    DOC.add_page_break()
    front_heading('MỤC LỤC')
    p=DOC.add_paragraph(); p.paragraph_format.first_line_indent=Cm(0)
    field(p,'TOC \\o "1-3" \\h \\z \\u','Cập nhật mục lục bằng Ctrl+A và F9 trong Word.')
    DOC.add_page_break()
    front_heading('DANH MỤC HÌNH')
    p=DOC.add_paragraph(); p.paragraph_format.first_line_indent=Cm(0)
    field(p,'TOC \\t "Chu thich hinh,1" \\h \\z','Cập nhật danh mục hình trong Word.')
    DOC.add_page_break()
    front_heading('DANH MỤC BẢNG')
    p=DOC.add_paragraph(); p.paragraph_format.first_line_indent=Cm(0)
    field(p,'TOC \\t "Chu thich bang,1" \\h \\z','Cập nhật danh mục bảng trong Word.')
    DOC.add_page_break()
    front_heading('DANH MỤC CHỮ VIẾT TẮT')


def table_title(header, last_heading):
    joined=' '.join(header).lower()
    if 'chữ viết tắt' in joined: return 'Danh mục chữ viết tắt'
    if 'mục tiêu cụ thể' in joined: return 'Mục tiêu và phương pháp đánh giá'
    if 'giá trị đối với bài toán' in joined: return 'Các hướng tiếp cận bài toán quản lý thiết bị'
    if 'trách nhiệm' in joined and 'tác nhân' in joined: return 'Tác nhân và giới hạn trách nhiệm'
    if 'yêu cầu' in joined and 'điều kiện chấp nhận' in joined: return 'Yêu cầu chức năng và điều kiện chấp nhận'
    if 'nội dung' in joined and 'cách đánh giá' in joined: return 'Yêu cầu phi chức năng và cách đánh giá'
    if 'hoạt động' in joined and 'employee' in joined: return 'Ma trận quyền nghiệp vụ'
    if 'thành phần chính' in joined: return 'Liên kết yêu cầu với thành phần và kiểm thử'
    if 'module' in joined: return 'Trách nhiệm và hợp đồng của các module'
    if 'thực thể khái niệm' in joined: return 'Các thực thể trong mô hình dữ liệu'
    if 'nguy cơ' in joined: return 'Nguy cơ và biện pháp kiểm soát'
    if 'khai báo' in joined: return 'Công nghệ và phiên bản khai báo trong dự án'
    if 'thông tin triển khai' in joined: return 'Thông tin môi trường triển khai'
    if 'cấu hình cần ghi' in joined: return 'Cấu hình môi trường kiểm thử'
    if 'nhóm kiểm tra' in joined: return 'Các nhóm kiểm tra và cách thu thập kết quả'
    if 'thao tác kiểm thử' in joined: return 'Kịch bản kiểm thử chức năng và quyền truy cập'
    if 'mẫu phát sinh' in joined: return 'Kết quả thực nghiệm mất mạng và gửi bù'
    if 'api p50/p95' in joined: return 'Kết quả hiệu năng theo mức tải'
    if 'chế độ agent' in joined: return 'Tài nguyên Agent theo chế độ hoạt động'
    if 'tiêu chí' in joined and 'bằng chứng' in joined: return 'Kết quả triển khai và phục hồi'
    if 'người tham gia' in joined: return 'Kết quả thực hiện tác vụ của người dùng thử'
    if 'bước' in joined and 'trình bày' in joined: return 'Trình tự demo và bằng chứng cần thấy'
    if 'trạng thái sau' in joined: return 'Đánh giá kết quả theo mục tiêu'
    if 'ưu tiên' in joined: return 'Ưu tiên phát triển và căn cứ thực hiện'
    if 'phương thức' in joined: return 'Các API tiêu biểu'
    if 'tên bảng vật lý' in joined: return 'Mẫu từ điển dữ liệu'
    if 'giai đoạn' in joined: return 'Giai đoạn thực hiện và điều kiện hoàn thành'
    if 'rủi ro' in joined: return 'Rủi ro và phương án xử lý'
    return last_heading or 'Thông tin tổng hợp'


def widths_for(header):
    n=len(header)
    h=' '.join(header)
    if n==2: return [5,11]
    if n==3:
        if header[0]=='Mã': return [1.5,7,7.5]
        return [4,6,6]
    if n==4:
        if header[0]=='Mã': return [1.65,5.6,5.45,3.3]
        if header[0]=='Mục tiêu': return [4,4.1,3.3,4.6]
        return [3.5,4,4.5,4]
    if n==5: return [4.4,3.2,3.1,3.1,2.2]
    if n==6: return [2.7,2.5,2.7,1.5,3.2,3.4]
    if n==7: return [1.8,2.2,2.3,2.4,2.5,2.2,2.6]
    return [16/n]*n


def add_table(rows, last_heading):
    global CURRENT_CHAPTER
    header=rows[0]
    if header[0]=='Chữ viết tắt':
        label='Bảng Danh mục chữ viết tắt'
    else:
        TABLE_COUNTS[CURRENT_CHAPTER]+=1
        c= str(CURRENT_CHAPTER) if CURRENT_CHAPTER else 'P'
        label=f'Bảng {c}.{TABLE_COUNTS[CURRENT_CHAPTER]} {table_title(header,last_heading)}'
    if header[0]!='Chữ viết tắt':
        p=DOC.add_paragraph(style='Chu thich bang')
        p.add_run(f'Bảng {c}.')
        field(p,f'SEQ Bang{c} \\* ARABIC',str(TABLE_COUNTS[CURRENT_CHAPTER]))
        p.add_run(' '+table_title(header,last_heading))
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        TABLES.append(label)
    t=DOC.add_table(rows=1,cols=len(header))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    widths=widths_for(header)
    for i,w in enumerate(widths): t.columns[i].width=Cm(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+name); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side,val in [('top',90),('bottom',90),('left',100),('right',100)]:
        a=OxmlElement('w:'+side); a.set(qn('w:w'),str(val)); a.set(qn('w:type'),'dxa'); margins.append(a)
    pr.append(margins)
    for ri,row in enumerate(rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        for ci,text in enumerate(row):
            c=cells[ci]; c.width=Cm(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0]; p.style=DOC.styles['Table Text']
            p.paragraph_format.keep_with_next=False
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.space_before=Pt(0)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ci==0 and widths[ci]<=3 else WD_ALIGN_PARAGRAPH.LEFT
            rich(p,text)
            if ri==0:
                p.paragraph_format.keep_with_next=True
                for r in p.runs: r.bold=True; r.font.color.rgb=BLACK
                shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'DDE6ED'); c._tc.get_or_add_tcPr().append(shd)
        trpr=t.rows[ri]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit'); trpr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); trpr.append(repeat)
    p=DOC.add_paragraph(); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1; p.paragraph_format.first_line_indent=Cm(0)


def strlist(x):
    if not isinstance(x,list): return str(x)
    parts=[str(v).strip() for v in x]
    separator=' ' if all(v.endswith(('.', '!', '?')) for v in parts) else '; '
    return separator.join(parts)


def add_figure(fid):
    if fid in FIGURE_INSERTED: return
    item=FIG[fid]; FIGURE_INSERTED.append(fid)
    caption=item['caption'].replace(' — ', ' ').replace(' – ', ' ')
    label={'diagram':'VẼ SƠ ĐỒ','screenshot':'CHỤP ẢNH','chart':'TẠO BIỂU ĐỒ'}.get(item['kind'],'CHÈN HÌNH')
    if fid=='H03' and (ROOT/'architecture.png').exists():
        p=DOC.add_paragraph(); p.paragraph_format.first_line_indent=Cm(0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(ROOT/'architecture.png'),width=Cm(16))
    else:
        note(f'{label} {fid} — {caption}.',keep=True)
        note('Cách thực hiện: '+strlist(item['capture_steps']),keep=True)
        note('Trong hình cần thấy: '+strlist(item['must_show']),keep=True)
        if item.get('hide'):
            note('Cần che hoặc bỏ: '+strlist(item['hide']),keep=True)
        p=DOC.add_paragraph('Chèn hình tại vị trí này',style='Normal')
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.keep_with_next=True
        for r in p.runs: r.italic=True; r.font.size=Pt(11.5)
    p=DOC.add_paragraph(style='Chu thich hinh')
    m=re.match(r'Hình (\d+)\.(\d+)\.\s*(.*)',caption)
    if m:
        p.add_run('Hình '+m.group(1)+'.')
        field(p,'SEQ Hinh'+m.group(1)+' \\* ARABIC',m.group(2))
        p.add_run('. '+m.group(3))
    else:
        p.add_run(caption)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    bookmark(p,fid)
    if fid!='H03':
        note('Gợi ý đoạn phân tích sau khi chèn hình: '+str(item['explanation']))


def apply_edits(text):
    path=ROOT/'editor-review.json'
    if path.exists():
        data=json.loads(path.read_text(encoding='utf-8'))
        items=data if isinstance(data,list) else data.get('edits',data.get('changes',[]))
        for item in items:
            original=item.get('original',''); replacement=item.get('replacement','')
            if original and original in text: text=text.replace(original,replacement)
    replacements={
        'Chuẩn hóa trích dẫn, tạo mục lục tự động':'Chuẩn hóa trích dẫn, cập nhật mục lục tự động',
        'trong mã nguồn đã đối chiếu':'trong mã nguồn của hệ thống',
        'Trong mã nguồn đã đối chiếu':'Trong mã nguồn của hệ thống',
        'trong phần mã đã đối chiếu':'trong phần xử lý tương ứng',
        'Phần mã đã đối chiếu cho thấy':'Hiện thực Agent cho thấy',
        'phần mã đã đối chiếu cho thấy':'hiện thực Agent cho thấy',
        'phần Core đã đối chiếu':'phần Core của Agent',
        'Trong phần Core đã đối chiếu':'Trong phần Core của Agent',
        'đã đối chiếu':'của hệ thống',
        'đã đọc':'của hệ thống',
        'đã kiểm tra':'được triển khai',
    }
    for a,b in replacements.items(): text=text.replace(a,b)
    return text


def parse(text):
    global CURRENT_CHAPTER
    text=apply_edits(text)
    # The existing cover is rebuilt; editorial guidance is presented once.
    start=text.index('## Lời cam đoan')
    text=text[start:]
    lines=text.splitlines()
    i=0; last_heading=''; in_references=False
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line.startswith('```'):
            i+=1
            while i<len(lines) and not lines[i].strip().startswith('```'): i+=1
            i+=1; continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{3,}:?',c or '') for c in row): rows.append(row)
                i+=1
            add_table(rows,last_heading); continue
        m=re.match(r'^(#{1,3}) (.+)$',line)
        if m:
            level=len(m.group(1)); title=m.group(2); last_heading=title
            if title=='Mục lục và các danh mục':
                lists(); i+=1
                while i<len(lines) and not lines[i].strip().startswith('|'): i+=1
                continue
            if level==1:
                if not title.startswith('PHỤ LỤC E '): DOC.add_page_break()
                ch=re.match(r'CHƯƠNG (\d+)',title)
                if ch: CURRENT_CHAPTER=int(ch.group(1))
                elif title.startswith('PHỤ LỤC'): CURRENT_CHAPTER=0
                hp=DOC.add_heading(title,level=1)
                if title.startswith('PHỤ LỤC E '): hp.paragraph_format.page_break_before=True
                in_references=title=='TÀI LIỆU THAM KHẢO'
            elif not CURRENT_CHAPTER and title in ['Lời cam đoan','Lời cảm ơn','Tóm tắt','Abstract']:
                if title!='Lời cam đoan': DOC.add_page_break()
                front_heading(title.upper())
            else:
                DOC.add_heading(title,level=level)
            i+=1; continue
        if line.startswith('[CHÈN HÌNH'):
            ids=list(dict.fromkeys(re.findall(r'\bH\d{2}\b',line)))
            for fid in ids:
                if fid in FIG: add_figure(fid)
            i+=1; continue
        if line.startswith('Hình 4.1 mô tả'):
            line=line.replace('Hình 4.1','Hình 4.1')
        if line.startswith(('[TỰ ĐIỀN','[XÁC NHẬN','[BỔ SUNG NGUỒN','[CHỤP HÌNH','[GHI CHÚ BIÊN TẬP')):
            note(line)
        elif line.startswith('[ĐO VÀ ĐIỀN'):
            note(line)
        elif re.match(r'^\d+\. ',line):
            p=para(line); p.paragraph_format.first_line_indent=Cm(0); p.paragraph_format.left_indent=Cm(0.4)
        elif line.startswith('- '):
            p=para('• '+line[2:]); p.paragraph_format.first_line_indent=Cm(-0.35); p.paragraph_format.left_indent=Cm(0.7)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        else:
            if in_references:
                p=para(line); p.paragraph_format.first_line_indent=Cm(-0.7); p.paragraph_format.left_indent=Cm(0.7)
            else:
                para(line)
        i+=1


def appendix_figure_guide():
    DOC.add_page_break()
    DOC.add_heading('PHỤ LỤC G HƯỚNG DẪN HÌNH MINH CHỨNG VÀ CHÚ THÍCH',level=1)
    para('Danh sách này tổng hợp vị trí hình để chuẩn bị minh chứng. Hướng dẫn chi tiết cũng được đặt ngay trong nội dung tại từng vị trí. Chú thích có thể giữ trong bản nộp; phần hướng dẫn chụp, vẽ và biên tập cần được bỏ sau khi hoàn thành.')
    for x in FIGURES:
        DOC.add_paragraph(x['caption'].replace(' — ',' ').replace(' – ',' '),style='Figure Guide Heading')
        note(f"Vị trí: mục {x['section']}; mã {x['id']}; loại {'sơ đồ' if x['kind']=='diagram' else 'ảnh chụp màn hình' if x['kind']=='screenshot' else 'biểu đồ từ dữ liệu đo'}.")
        if x['id']=='H03':
            note('Sơ đồ kiến trúc tổng thể đã được chèn sẵn. Kiểm tra lại các thành phần trước khi nộp.')
        else:
            note('Cách thực hiện: '+strlist(x['capture_steps']))
            note('Nội dung cần nhìn thấy: '+strlist(x['must_show']))
            if x.get('hide'): note('Thông tin cần che: '+strlist(x['hide']))
        para('Chú thích đề xuất: '+x['caption'])


setup()
cover()
editor_guide()
parse(SOURCE.read_text(encoding='utf-8'))
appendix_figure_guide()
assert sorted(FIGURE_INSERTED)==sorted(FIG), (FIGURE_INSERTED, list(FIG))
DOC.save(OUT)
(ROOT/'word-build-summary.json').write_text(json.dumps({'output':str(OUT),'paragraphs':len(DOC.paragraphs),'tables':len(DOC.tables),'figures':FIGURE_INSERTED,'table_captions':TABLES},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(OUT),'paragraphs':len(DOC.paragraphs),'tables':len(DOC.tables),'figures':len(FIGURE_INSERTED)},ensure_ascii=False))
